from calendar import monthrange
from collections import defaultdict
from datetime import date
from io import BytesIO

from django.db.models import OuterRef, Subquery, Sum
from django.db.models.functions import TruncDate
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.views import APIView
from rest_framework.response import Response

from usuarios.permissions import DashboardPermission, RolePermission
from auditoria.services import record_action
from consumos.models import Consumo
from dispositivos.models import Dispositivo
from sims.models import Sim
from bateria.models import RegistroBateria

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(fill_type='solid', fgColor='007F5F')
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _month_range(year, month):
    start = date(year, month, 1)
    end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    return start, end


def _last_day(year, month):
    """Mes actual: hasta hoy. Mes pasado: completo. Mes futuro: ninguno."""
    today = timezone.localdate()
    if (year, month) == (today.year, today.month):
        return today.day
    if (year, month) > (today.year, today.month):
        return 0
    return monthrange(year, month)[1]


def _style_header(row_cells):
    for cell in row_cells:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


class DashboardSummaryView(APIView):
    permission_classes = [DashboardPermission]

    def get(self, request):
        try:
            year = int(request.query_params.get('anio', timezone.localdate().year))
            month = int(request.query_params.get('mes', timezone.localdate().month))
        except (TypeError, ValueError):
            return Response({'detail': 'El año y el mes deben ser números.'}, status=400)
        if month < 1 or month > 12:
            return Response({'detail': 'El mes debe estar entre 1 y 12.'}, status=400)

        monthly = Consumo.objects.filter(periodo='diario', fecha__year=year, fecha__month=month)
        consumption = monthly.annotate(
            day=TruncDate('fecha', tzinfo=timezone.get_current_timezone())
        ).values('day', 'sim_id').annotate(
            consumo_datos_movil=Sum('consumo_datos_movil'),
        ).order_by('day', 'sim_id')

        latest_battery = RegistroBateria.objects.filter(
            dispositivo_id=OuterRef('pk')
        ).order_by('-fecha_hora')
        devices = Dispositivo.objects.annotate(
            battery_percent=Subquery(latest_battery.values('porcentaje')[:1]),
            battery_state=Subquery(latest_battery.values('estado')[:1]),
        ).values('id', 'modelo', 'device_uuid', 'activo', 'ultimo_contacto', 'battery_percent', 'battery_state')
        sims = Sim.objects.select_related('operador').values('id', 'numero_telefonico', 'operador__nombre')

        return Response({
            'metrics': {
                'devices': Dispositivo.objects.count(),
                'sims': Sim.objects.count(),
                'connected': Dispositivo.objects.filter(activo=True).count(),
            },
            'devices': list(devices),
            'sims': [
                {
                    'id': sim['id'],
                    'numero_telefonico': sim['numero_telefonico'],
                    'operador_nombre': sim['operador__nombre'],
                }
                for sim in sims
            ],
            'consumption': [
                {
                    'fecha': row['day'].isoformat(),
                    'periodo': 'diario',
                    'sim': row['sim_id'],
                    'consumo_datos_movil': row['consumo_datos_movil'] or 0,
                }
                for row in consumption
            ],
        })


class SimMonthlyExportView(APIView):
    permission_classes = [RolePermission]

    def get(self, request, pk):
        month = int(request.query_params.get('mes', timezone.now().month))
        year = int(request.query_params.get('anio', timezone.now().year))
        if month < 1 or month > 12:
            return HttpResponse('El mes debe estar entre 1 y 12.', status=400)

        sim = Sim.objects.select_related('dispositivo', 'operador').filter(pk=pk).first()
        if not sim:
            return HttpResponse('SIM no encontrada.', status=404)

        start, end = _month_range(year, month)
        consumos = Consumo.objects.filter(
            sim=sim, periodo='diario', fecha__date__gte=start, fecha__date__lt=end
        )

        por_dia = defaultdict(float)
        for item in consumos:
            por_dia[timezone.localtime(item.fecha).date()] += item.consumo_datos_movil
        total_mes = round(sum(por_dia.values()), 2)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Resumen SIM'

        # Resumen del mes
        sheet.append(['Fecha', 'Serial dispositivo', 'Datos móviles usados en el mes (MB)', 'Operador', 'ICCID', 'Número telefónico'])
        _style_header(sheet[sheet.max_row])
        sheet.append([
            f'{year}-{month:02d}',
            sim.dispositivo.serial if sim.dispositivo else '',
            total_mes,
            sim.operador.nombre if sim.operador else '',
            sim.iccid or '',
            sim.numero_telefonico or '',
        ])

        # Detalle día por día
        sheet.append([])
        sheet.append(['Día', 'Datos móviles usados en el día (MB)', 'Acumulado del mes (MB)'])
        _style_header(sheet[sheet.max_row])
        acumulado = 0
        for day in range(1, _last_day(year, month) + 1):
            current = date(year, month, day)
            usado = por_dia.get(current)
            if usado is not None:
                acumulado += usado
            sheet.append([current, round(usado, 2) if usado is not None else 'Sin registro', round(acumulado, 2)])
            sheet.cell(row=sheet.max_row, column=1).number_format = 'yyyy-mm-dd'
        sheet.append(['TOTAL', total_mes, ''])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

        for column, width in {'A': 26, 'B': 38, 'C': 38, 'D': 26, 'E': 26, 'F': 26}.items():
            sheet.column_dimensions[column].width = width

        output = BytesIO()
        workbook.save(output)
        record_action(request, 'EXPORTAR', 'consumos', sim.pk, f'Reporte mensual SIM {year}-{month:02d}')
        response = HttpResponse(output.getvalue(), content_type=XLSX_TYPE)
        response['Content-Disposition'] = f'attachment; filename="sim-{sim.pk}-{year}-{month:02d}.xlsx"'
        return response


class DevicesMonthlyExportView(APIView):
    permission_classes = [RolePermission]

    def get(self, request):
        month = int(request.query_params.get('mes', timezone.now().month))
        year = int(request.query_params.get('anio', timezone.now().year))
        if month < 1 or month > 12:
            return HttpResponse('El mes debe estar entre 1 y 12.', status=400)

        start, end = _month_range(year, month)

        por_dia = defaultdict(float)
        total_por_dispositivo = defaultdict(float)
        consumos = Consumo.objects.filter(periodo='diario', fecha__date__gte=start, fecha__date__lt=end)
        for item in consumos:
            dia = timezone.localtime(item.fecha).date()
            por_dia[(item.dispositivo_id, dia)] += item.consumo_datos_movil
            total_por_dispositivo[item.dispositivo_id] += item.consumo_datos_movil

        dispositivos = list(Dispositivo.objects.all().order_by('id'))
        workbook = Workbook()

        # Hoja 1: resumen del mes por dispositivo
        sheet = workbook.active
        sheet.title = 'Dispositivos'
        sheet.append(['Fecha', 'Serial', 'Modelo', 'Fabricante', 'Estado', 'Datos móviles usados en el mes (MB)', 'ICCID', 'IMEI'])
        _style_header(sheet[1])

        iccids = {}
        for device in dispositivos:
            sim = device.sims.order_by('id').first()
            iccids[device.id] = sim.iccid if sim and sim.iccid else ''
            sheet.append([
                f'{year}-{month:02d}', device.serial or '', device.modelo or '', device.fabricante or '',
                'Activo' if device.activo else 'Inactivo',
                round(total_por_dispositivo.get(device.id, 0), 2),
                iccids[device.id], device.imei_1 or '',
            ])
        sheet.freeze_panes = 'A2'
        for column in ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H'):
            sheet.column_dimensions[column].width = 25
        sheet.column_dimensions['F'].width = 38

        # Hoja 2: día por día, por dispositivo
        daily = workbook.create_sheet('Consumo diario')
        daily.append(['Día', 'Serial', 'Modelo', 'ICCID', 'Datos móviles usados en el día (MB)'])
        _style_header(daily[1])
        for device in dispositivos:
            for day in range(1, _last_day(year, month) + 1):
                current = date(year, month, day)
                usado = por_dia.get((device.id, current))
                daily.append([
                    current, device.serial or '', device.modelo or '', iccids[device.id],
                    round(usado, 2) if usado is not None else 'Sin registro',
                ])
                daily.cell(row=daily.max_row, column=1).number_format = 'yyyy-mm-dd'
        daily.freeze_panes = 'A2'
        for column in ('A', 'B', 'C', 'D', 'E'):
            daily.column_dimensions[column].width = 26
        daily.column_dimensions['E'].width = 38
        output = BytesIO()
        workbook.save(output)
        record_action(request, 'EXPORTAR', 'dispositivos', '', f'Reporte mensual dispositivos {year}-{month:02d}')
        response = HttpResponse(output.getvalue(), content_type=XLSX_TYPE)
        response['Content-Disposition'] = f'attachment; filename="dispositivos-{year}-{month:02d}.xlsx"'
        return response